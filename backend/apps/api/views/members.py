import csv
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import status
from rest_framework.filters import OrderingFilter, SearchFilter
from rest_framework.generics import ListAPIView
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.api.pagination import StandardPagination
from apps.api.permissions import (
    CanExportMembers,
    CanManageMembers,
    CanValidateMembership,
    CanViewMembers,
)
from apps.api.serializers.members import (
    AdminMemberDetailSerializer,
    AdminMemberListSerializer,
    AdminMemberStatusSerializer,
    AdminMemberUpdateSerializer,
)
from apps.core.mixins import AuditMixin
from apps.core.models import AuditLog
from apps.members.models import MemberProfile


class AdminMemberListView(ListAPIView):
    """Admin : liste paginee des membres avec recherche et filtres."""

    permission_classes = [IsAuthenticated, CanViewMembers]
    serializer_class = AdminMemberListSerializer
    pagination_class = StandardPagination
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["membership_status", "city", "commune", "registration_source"]
    search_fields = [
        "matricule",
        "user__first_name", "user__last_name",
        "user__email", "user__phone",
        "city", "commune",
    ]
    ordering_fields = ["created_at", "city", "membership_status"]
    ordering = ["-created_at"]

    def get_queryset(self):
        qs = MemberProfile.objects.select_related("user")

        sex = self.request.query_params.get("sex")
        if sex in ("M", "F"):
            qs = qs.filter(user__sex=sex)

        is_active = self.request.query_params.get("is_active")
        if is_active == "true":
            qs = qs.filter(user__is_active=True)
        elif is_active == "false":
            qs = qs.filter(user__is_active=False)

        return qs


class AdminMemberDetailView(APIView):
    """Admin : fiche detaillee d'un membre."""

    permission_classes = [IsAuthenticated, CanViewMembers]

    def get(self, request, pk):
        try:
            profile = (
                MemberProfile.objects
                .select_related("user")
                .prefetch_related("user__party_roles__role", "user__party_roles__zone")
                .get(pk=pk)
            )
        except MemberProfile.DoesNotExist:
            return Response({"detail": "Membre introuvable."}, status=status.HTTP_404_NOT_FOUND)

        return Response(AdminMemberDetailSerializer(profile).data)


class AdminMemberUpdateView(APIView):
    """Admin : modifier les informations d'un membre."""

    permission_classes = [IsAuthenticated, CanManageMembers]
    parser_classes = [MultiPartParser, FormParser]

    def patch(self, request, pk):
        try:
            profile = MemberProfile.objects.select_related("user").get(pk=pk)
        except MemberProfile.DoesNotExist:
            return Response({"detail": "Membre introuvable."}, status=status.HTTP_404_NOT_FOUND)

        old_values = {
            field: getattr(profile, field)
            for field in request.data.keys()
            if hasattr(profile, field)
        }

        serializer = AdminMemberUpdateSerializer(profile, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        instance = serializer.save()

        clear_photo = request.data.get("clear_photo", "").lower() in ("true", "1")
        if clear_photo and not request.FILES.get("photo"):
            if instance.photo:
                instance.photo.delete(save=False)
            instance.photo = None
            instance.save(update_fields=["photo"])

        AuditMixin.log_action(
            user=request.user,
            action=AuditLog.ActionChoices.UPDATE,
            entity_type="MemberProfile",
            entity_id=profile.pk,
            changes={k: {"old": str(v), "new": str(getattr(profile, k))} for k, v in old_values.items()},
            request=request,
        )

        profile.refresh_from_db()
        return Response(AdminMemberDetailSerializer(
            MemberProfile.objects
            .select_related("user")
            .prefetch_related("user__party_roles__role", "user__party_roles__zone")
            .get(pk=pk)
        ).data)


class AdminMemberStatusView(APIView):
    """Admin : changer le statut d'adhesion d'un membre."""

    permission_classes = [IsAuthenticated, CanValidateMembership]

    def patch(self, request, pk):
        try:
            profile = MemberProfile.objects.select_related("user").get(pk=pk)
        except MemberProfile.DoesNotExist:
            return Response({"detail": "Membre introuvable."}, status=status.HTTP_404_NOT_FOUND)

        serializer = AdminMemberStatusSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        new_status = serializer.validated_data["status"]
        old_status = profile.membership_status

        if old_status == new_status:
            return Response(
                {"detail": f"Le membre est déjà '{profile.get_membership_status_display()}'."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        with transaction.atomic():
            profile.membership_status = new_status

            if new_status == MemberProfile.StatusChoices.VALIDATED and not profile.matricule:
                from apps.members.utils import generate_hmac_matricule
                profile.matricule = generate_hmac_matricule(profile.user, profile.city)
                profile.membership_date = timezone.now()
                profile.save(update_fields=[
                    "membership_status", "matricule", "membership_date", "updated_at",
                ])
            else:
                profile.save(update_fields=["membership_status", "updated_at"])

            AuditMixin.log_action(
                user=request.user,
                action=AuditLog.ActionChoices.UPDATE,
                entity_type="MemberProfile",
                entity_id=profile.pk,
                changes={
                    "membership_status": {"old": old_status, "new": new_status},
                    "reason": serializer.validated_data.get("reason", ""),
                },
                request=request,
            )

        return Response({
            "detail": f"Statut modifié : {profile.get_membership_status_display()}.",
            "membership_status": profile.membership_status,
            "matricule": profile.matricule,
        })


class AdminMemberExportView(APIView):
    """Admin : export Excel/CSV des membres (filtres appliques)."""

    permission_classes = [IsAuthenticated, CanExportMembers]

    def get(self, request):
        qs = MemberProfile.objects.select_related("user").order_by("-created_at")

        membership_status = request.query_params.get("membership_status")
        if membership_status:
            qs = qs.filter(membership_status=membership_status)
        city = request.query_params.get("city")
        if city:
            qs = qs.filter(city__icontains=city)
        commune = request.query_params.get("commune")
        if commune:
            qs = qs.filter(commune__icontains=commune)
        sex = request.query_params.get("sex")
        if sex in ("M", "F"):
            qs = qs.filter(user__sex=sex)

        export_format = request.query_params.get("format", "csv")

        if export_format == "excel":
            return self._export_excel(request, qs)
        return self._export_csv(request, qs)

    def _export_csv(self, request, queryset):
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="membres_fpp.csv"'
        response.write("\ufeff")

        writer = csv.writer(response, delimiter=";")
        writer.writerow([
            "Matricule", "Nom", "Prénom", "Email", "Téléphone", "Sexe",
            "Ville", "Commune", "Région", "Quartier",
            "Profession", "Statut", "Source", "Date adhésion", "Date inscription",
        ])

        for p in queryset.iterator(chunk_size=500):
            writer.writerow([
                p.matricule or "",
                p.user.last_name,
                p.user.first_name,
                p.user.email,
                p.user.phone or "",
                p.user.get_sex_display() if p.user.sex else "",
                p.city,
                p.commune,
                p.region,
                p.neighborhood,
                p.profession,
                p.get_membership_status_display(),
                p.get_registration_source_display(),
                p.membership_date.strftime("%d/%m/%Y") if p.membership_date else "",
                p.created_at.strftime("%d/%m/%Y"),
            ])

        self._log_export(request, queryset)
        return response

    def _export_excel(self, request, queryset):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Membres FPP"

        headers = [
            "Matricule", "Nom", "Prénom", "Email", "Téléphone", "Sexe",
            "Ville", "Commune", "Région", "Quartier",
            "Profession", "Statut", "Source", "Date adhésion", "Date inscription",
        ]
        header_font = Font(bold=True)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_idx, p in enumerate(queryset.iterator(chunk_size=500), 2):
            ws.cell(row=row_idx, column=1, value=p.matricule or "")
            ws.cell(row=row_idx, column=2, value=p.user.last_name)
            ws.cell(row=row_idx, column=3, value=p.user.first_name)
            ws.cell(row=row_idx, column=4, value=p.user.email)
            ws.cell(row=row_idx, column=5, value=p.user.phone or "")
            ws.cell(row=row_idx, column=6, value=p.user.get_sex_display() if p.user.sex else "")
            ws.cell(row=row_idx, column=7, value=p.city)
            ws.cell(row=row_idx, column=8, value=p.commune)
            ws.cell(row=row_idx, column=9, value=p.region)
            ws.cell(row=row_idx, column=10, value=p.neighborhood)
            ws.cell(row=row_idx, column=11, value=p.profession)
            ws.cell(row=row_idx, column=12, value=p.get_membership_status_display())
            ws.cell(row=row_idx, column=13, value=p.get_registration_source_display())
            ws.cell(row=row_idx, column=14, value=p.membership_date.strftime("%d/%m/%Y") if p.membership_date else "")
            ws.cell(row=row_idx, column=15, value=p.created_at.strftime("%d/%m/%Y"))

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = 18

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="membres_fpp.xlsx"'

        self._log_export(request, queryset)
        return response

    def _log_export(self, request, queryset):
        AuditMixin.log_action(
            user=request.user,
            action=AuditLog.ActionChoices.EXPORT,
            entity_type="MemberProfile",
            entity_id=request.user.pk,
            changes={"count": queryset.count(), "filters": dict(request.query_params)},
            request=request,
        )
